from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


COLORS = {
    "paper": "FBF7EF",
    "surface": "FFFDF8",
    "wash": "F7EFE4",
    "line": "DED2C2",
    "line_strong": "C7B39E",
    "brand": "A94F2D",
    "brand_dark": "71361F",
    "green": "31594D",
    "green_soft": "E8F3ED",
    "best_cell": "EDF6F0",
    "ink": "211B16",
    "muted": "6A6056",
    "white": "FFFFFF",
}


FONT = "Aptos"
THIN = Side(style="thin", color=COLORS["line"])
MEDIUM_GREEN = Side(style="medium", color=COLORS["green"])
MEDIUM_LINE = Side(style="medium", color=COLORS["line_strong"])


def money(value: float | None) -> str:
    if value is None:
        return "-"
    text = f"R$ {value:,.2f}"
    return text.replace(",", "X").replace(".", ",").replace("X", ".")


def pct(value: float | None) -> str:
    if value is None:
        return "-"
    return f"{value:.1f}%"


def fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)


def set_border(cell, *, color: str = COLORS["line"], style: str = "thin") -> None:
    side = Side(style=style, color=color)
    cell.border = Border(left=side, right=side, top=side, bottom=side)


def style_range(ws, cell_range: str, *, fill_color=None, font=None, border=True, alignment=None):
    for row in ws[cell_range]:
        for cell in row:
            if fill_color:
                cell.fill = fill(fill_color)
            if font:
                cell.font = font
            if border:
                set_border(cell)
            if alignment:
                cell.alignment = alignment


def merge_label(ws, cell_range: str, value: str, *, fill_color: str, font_size=14):
    ws.merge_cells(cell_range)
    cell = ws[cell_range.split(":")[0]]
    cell.value = value
    cell.fill = fill(fill_color)
    cell.font = Font(name=FONT, bold=True, size=font_size, color=COLORS["ink"])
    cell.alignment = Alignment(vertical="center")


def setup_sheet(ws, title: str):
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A1"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.35
    ws.page_margins.bottom = 0.35
    ws["A1"] = title
    ws["A1"].font = Font(name=FONT, bold=True, size=16, color=COLORS["ink"])
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 28


def write_dashboard(wb: Workbook, data: dict):
    ws = wb.create_sheet("Dashboard")
    customer_name = data.get("customer", {}).get("name", "Cliente")
    number = data.get("number") or data.get("pedido", {}).get("number") or ""
    quote_type = data.get("type", "")
    setup_sheet(ws, f"Mapa de Cotação | Pedido {number} {quote_type} | {customer_name}")
    ws.freeze_panes = "A5"

    for col, width in {
        "A": 22, "B": 20, "C": 20, "D": 20, "E": 20, "F": 20, "G": 20, "H": 20,
    }.items():
        ws.column_dimensions[col].width = width

    ws.merge_cells("A2:H2")
    ws["A2"] = "Resumo executivo gerado a partir do JSON único do mapa de cotação"
    ws["A2"].font = Font(name=FONT, size=10, color=COLORS["muted"])

    kpis = [
        ("CESTA SUGERIDA", money(data["summary"]["basketTotal"]), "menor preço por item cotado"),
        ("ECONOMIA POSSÍVEL", money(data["summary"]["basketSavings"]), "comparado ao pior conjunto cotado"),
        ("ITENS SEM PREÇO", data["summary"]["itemsWithoutQuotes"], "precisam retorno antes do pedido"),
        ("ITENS RESOLVIDOS", data["summary"]["basketResolved"], "já entram na decisão de compra"),
        ("FORNECEDORES", data["summary"]["supplierCount"], "comparados neste mapa"),
        ("GRUPOS", data["summary"]["groupCount"], "famílias técnicas para conferência"),
    ]

    start_col = 1
    for idx, (label, value, note) in enumerate(kpis):
        col = start_col + idx
        cell = ws.cell(4, col)
        cell.value = f"{label}\n{value}\n{note}"
        cell.fill = fill(COLORS["surface"])
        cell.font = Font(name=FONT, bold=True, size=11, color=COLORS["ink"])
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        set_border(cell, color=COLORS["line_strong"])
    ws.row_dimensions[4].height = 74

    row = 7
    merge_label(ws, f"A{row}:H{row}", "Navegação por grupos", fill_color=COLORS["wash"], font_size=13)
    row += 1
    for group_index, group in enumerate(data["groups"], start=1):
        ws.cell(row, 1).value = group["title"]
        ws.cell(row, 2).value = f"{group['resolved']}/{len(group['items'])}"
        ws.cell(row, 3).value = money(group.get("bestOfferTotal"))
        ws.cell(row, 4).value = money(group.get("savings"))
        for col in range(1, 5):
            ws.cell(row, col).font = Font(name=FONT, bold=col == 1, size=10, color=COLORS["ink"])
            ws.cell(row, col).alignment = Alignment(vertical="center")
            set_border(ws.cell(row, col))
        ws.cell(row, 1).hyperlink = f"#'Mapa de Cotacao'!A{group['_excel_row']}"
        ws.cell(row, 1).style = "Hyperlink"
        row += 1
    row += 1
    ws.cell(row, 1).value = "Itens faltantes"
    ws.cell(row, 2).value = len(data.get("pendingItems", []))
    ws.cell(row, 1).hyperlink = f"#'Mapa de Cotacao'!A{data['_pending_row']}"
    ws.cell(row, 1).style = "Hyperlink"
    for col in range(1, 5):
        set_border(ws.cell(row, col))

    row += 3
    merge_label(ws, f"A{row}:H{row}", "Distribuição da cesta sugerida por fornecedor", fill_color=COLORS["wash"], font_size=13)
    row += 1
    headers = ["Fornecedor", "Total", "% da cesta", "Itens"]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row, col)
        cell.value = header
        cell.fill = fill(COLORS["wash"])
        cell.font = Font(name=FONT, bold=True, size=10, color=COLORS["muted"])
        set_border(cell)
    row += 1
    total = data["summary"]["basketTotal"] or 1
    for allocation in data["basket"]["supplierAllocations"]:
        values = [
            allocation["name"],
            allocation["total"],
            allocation["total"] / total,
            allocation["count"],
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row, col)
            cell.value = value
            cell.font = Font(name=FONT, bold=col in (1, 2), size=10, color=COLORS["ink"])
            cell.alignment = Alignment(vertical="center")
            set_border(cell)
        ws.cell(row, 2).number_format = '"R$" #,##0.00'
        ws.cell(row, 3).number_format = "0.0%"
        row += 1

    row += 2
    merge_label(ws, f"A{row}:H{row}", "Pendências sem cotação", fill_color=COLORS["wash"], font_size=13)
    row += 1
    for col, header in enumerate(["Item", "Descrição", "Quantidade", "Unidade", "Grupo"], start=1):
        cell = ws.cell(row, col)
        cell.value = header
        cell.fill = fill(COLORS["wash"])
        cell.font = Font(name=FONT, bold=True, size=10, color=COLORS["muted"])
        set_border(cell)
    row += 1
    for item in data.get("pendingItems", []):
        values = [item["position"], item["description"], item["quantity"], item["unit"], item.get("title") or item.get("groupTitle") or "-"]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row, col)
            cell.value = value
            cell.font = Font(name=FONT, size=10, color=COLORS["ink"])
            cell.alignment = Alignment(wrap_text=col == 2, vertical="center")
            set_border(cell)
        row += 1


def write_map(ws, data: dict):
    setup_sheet(ws, "Mapa de Cotação | Matriz por grupos")
    ws.freeze_panes = "D4"

    suppliers = data["suppliers"]
    widths = [48, 10, 9] + [16] * len(suppliers)
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    row = 3
    for group in data["groups"]:
        group["_excel_row"] = row
        best_total_supplier_id = (group.get("bestTotalSupplier") or {}).get("supplierId")
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        ws.cell(row, 1).value = group["title"]
        ws.cell(row, 1).fill = fill(COLORS["wash"])
        ws.cell(row, 1).font = Font(name=FONT, bold=True, size=13, color=COLORS["ink"])
        ws.cell(row, 1).alignment = Alignment(vertical="center")
        metric_labels = [
            ("ITENS", len(group["items"])),
            ("COBERTOS", f"{group['resolved']}/{len(group['items'])}"),
            ("MELHOR GRUPO", money(group.get("bestOfferTotal"))),
            ("ECONOMIA", money(group.get("savings"))),
        ]
        metric_col = 4
        for label, value in metric_labels:
            cell = ws.cell(row, metric_col)
            cell.value = f"{label}\n{value}"
            cell.fill = fill(COLORS["wash"])
            cell.font = Font(name=FONT, bold=True, size=10, color=COLORS["ink"])
            cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
            metric_col += 1
        for col in range(1, 4 + len(suppliers)):
            set_border(ws.cell(row, col), color=COLORS["line"])
        ws.row_dimensions[row].height = 36
        row += 1

        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3 + len(suppliers))
        ws.cell(row, 1).value = group.get("subtitle") or "Família técnica de materiais"
        ws.cell(row, 1).fill = fill(COLORS["wash"])
        ws.cell(row, 1).font = Font(name=FONT, size=10, color=COLORS["muted"])
        ws.cell(row, 1).alignment = Alignment(vertical="center")
        row += 1

        headers = ["Item", "Qtd.", "Un."] + [supplier["sheetName"] for supplier in suppliers]
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row, col)
            cell.value = header
            cell.fill = fill(COLORS["wash"])
            cell.font = Font(name=FONT, bold=True, size=10, color=COLORS["muted"])
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            set_border(cell, color=COLORS["line_strong"])
            if col > 3 and suppliers[col - 4]["id"] == best_total_supplier_id:
                cell.value = f"{header}\nMENOR TOTAL"
                cell.fill = fill(COLORS["green_soft"])
                cell.font = Font(name=FONT, bold=True, size=10, color=COLORS["green"])
                cell.border = Border(left=MEDIUM_GREEN, right=MEDIUM_GREEN, top=MEDIUM_GREEN, bottom=THIN)
        ws.row_dimensions[row].height = 34
        header_row = row
        row += 1

        for item in group["items"]:
            ws.cell(row, 1).value = f"{item['position']}. {item['description']}"
            ws.cell(row, 2).value = item["quantity"]
            ws.cell(row, 3).value = item["unit"]
            for col in range(1, 4):
                ws.cell(row, col).font = Font(name=FONT, bold=col == 1, size=10, color=COLORS["ink"])
                ws.cell(row, col).alignment = Alignment(wrap_text=col == 1, vertical="center", horizontal="right" if col == 2 else "left")
                ws.cell(row, col).fill = fill(COLORS["white"])
                set_border(ws.cell(row, col))

            for supplier_index, supplier in enumerate(suppliers, start=4):
                offer = next((entry for entry in item.get("offers", []) if entry["supplierId"] == supplier["id"] and entry.get("available")), None)
                cell = ws.cell(row, supplier_index)
                if offer and offer.get("unitPrice") is not None:
                    cell.value = f"{money(offer['totalPrice'])}\n{money(offer['unitPrice'])}/un"
                else:
                    cell.value = "-"
                is_best_price = item.get("bestSupplierId") == supplier["id"]
                is_best_total = supplier["id"] == best_total_supplier_id
                cell.fill = fill(COLORS["best_cell"] if is_best_price else COLORS["white"])
                cell.font = Font(name=FONT, bold=is_best_price, size=9, color=COLORS["green"] if is_best_price else COLORS["ink"])
                cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="left")
                set_border(cell)
                if is_best_total:
                    cell.border = Border(left=MEDIUM_GREEN, right=MEDIUM_GREEN, top=THIN, bottom=THIN)
                if is_best_price:
                    cell.comment = Comment("Menor valor unitário para este item.", "Dema")
            ws.row_dimensions[row].height = 43
            row += 1

        ws.cell(row, 1).value = "TOTAL DO GRUPO POR FORNECEDOR"
        ws.cell(row, 1).font = Font(name=FONT, bold=True, size=10, color=COLORS["brand_dark"])
        ws.cell(row, 1).fill = fill(COLORS["wash"])
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        for col in range(1, 4):
            set_border(ws.cell(row, col), color=COLORS["line_strong"])
        for supplier_index, supplier in enumerate(suppliers, start=4):
            supplier_group = next((entry for entry in group.get("suppliers", []) if entry["supplierId"] == supplier["id"]), None)
            cell = ws.cell(row, supplier_index)
            if supplier_group and supplier_group.get("covered"):
                cell.value = f"{money(supplier_group['total'])}\n{supplier_group['covered']} itens cotados"
            else:
                cell.value = "-\nsem itens"
            is_best_total = supplier["id"] == best_total_supplier_id
            cell.fill = fill(COLORS["green_soft"] if is_best_total else COLORS["wash"])
            cell.font = Font(name=FONT, bold=True, size=9, color=COLORS["green"] if is_best_total else COLORS["ink"])
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            set_border(cell, color=COLORS["line_strong"])
            if is_best_total:
                cell.border = Border(left=MEDIUM_GREEN, right=MEDIUM_GREEN, top=THIN, bottom=MEDIUM_GREEN)
        ws.row_dimensions[row].height = 42
        row += 2

    data["_pending_row"] = row
    merge_label(ws, f"A{row}:{get_column_letter(3 + len(suppliers))}{row}", "Pendências sem cotação", fill_color=COLORS["wash"], font_size=13)
    row += 1
    for col, header in enumerate(["Item", "Qtd.", "Un.", "Descrição"], start=1):
        cell = ws.cell(row, col)
        cell.value = header
        cell.fill = fill(COLORS["wash"])
        cell.font = Font(name=FONT, bold=True, size=10, color=COLORS["muted"])
        set_border(cell)
    row += 1
    for item in data.get("pendingItems", []):
        values = [item["position"], item["quantity"], item["unit"], item["description"]]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row, col)
            cell.value = value
            cell.font = Font(name=FONT, size=10, color=COLORS["ink"])
            cell.alignment = Alignment(wrap_text=col == 4, vertical="center")
            set_border(cell)
        row += 1

    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(3 + len(suppliers))}{header_row}"


def write_data(wb: Workbook, data: dict):
    ws = wb.create_sheet("Dados")
    setup_sheet(ws, "Dados normalizados do JSON")
    ws.freeze_panes = "A2"
    headers = ["Grupo", "Item", "Descrição", "Qtd", "Un", "Fornecedor", "Total", "Unitário", "Melhor Item"]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(2, col)
        cell.value = header
        cell.fill = fill(COLORS["wash"])
        cell.font = Font(name=FONT, bold=True, size=10, color=COLORS["muted"])
        set_border(cell)
    row = 3
    for group in data["groups"]:
        for item in group["items"]:
            for offer in item.get("offers", []):
                values = [
                    group["title"],
                    item["position"],
                    item["description"],
                    item["quantity"],
                    item["unit"],
                    offer["sheet"],
                    offer.get("totalPrice") if offer.get("available") else None,
                    offer.get("unitPrice") if offer.get("available") else None,
                    offer.get("supplierId") == item.get("bestSupplierId"),
                ]
                for col, value in enumerate(values, start=1):
                    cell = ws.cell(row, col)
                    cell.value = value
                    cell.font = Font(name=FONT, size=9, color=COLORS["ink"])
                    cell.alignment = Alignment(wrap_text=col == 3, vertical="center")
                    set_border(cell)
                row += 1
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 8
    ws.column_dimensions["C"].width = 52
    ws.column_dimensions["D"].width = 10
    ws.column_dimensions["E"].width = 8
    ws.column_dimensions["F"].width = 18
    ws.column_dimensions["G"].width = 14
    ws.column_dimensions["H"].width = 14
    ws.column_dimensions["I"].width = 12
    for col in ("G", "H"):
        for cell in ws[col][2:]:
            cell.number_format = '"R$" #,##0.00'


def parse_args():
    parser = argparse.ArgumentParser(description="Build the Dema quotation-map XLSX dashboard.")
    parser.add_argument("input", type=Path, help="Canonical map JSON path.")
    parser.add_argument("output", type=Path, help="Output XLSX path.")
    return parser.parse_args()


def main():
    args = parse_args()
    data = json.loads(args.input.read_text(encoding="utf-8"))
    wb = Workbook()
    wb.remove(wb.active)

    map_ws = wb.create_sheet("Mapa de Cotacao")
    write_map(map_ws, data)
    write_dashboard(wb, data)
    write_data(wb, data)
    wb.active = wb.sheetnames.index("Dashboard")

    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    cell.alignment = Alignment(
                        horizontal=cell.alignment.horizontal,
                        vertical=cell.alignment.vertical or "center",
                        text_rotation=cell.alignment.text_rotation,
                        wrap_text=cell.alignment.wrap_text,
                        shrink_to_fit=cell.alignment.shrink_to_fit,
                        indent=cell.alignment.indent,
                    )

    args.output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.output)
    print(args.output)


if __name__ == "__main__":
    main()
