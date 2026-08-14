import json
from openpyxl import load_workbook
wb = load_workbook('.fractal/artifacts/docs/templates/PEDIDO-GERAL.xlsx')
ws = wb[wb.sheetnames[0]]
result = {
    "D14": ws['D14'].value,
    "D15": ws['D15'].value,
    "D16": ws['D16'].value,
    "D17": ws['D17'].value,
}
with open('.fractal/artifacts/docs/templates/check_result.json', 'w') as f:
    json.dump(result, f, indent=2)
