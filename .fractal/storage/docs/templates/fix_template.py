from openpyxl import load_workbook
wb = load_workbook('.fractal/artifacts/docs/templates/PEDIDO-GERAL.xlsx')
ws = wb[wb.sheetnames[0]]
ws['D15'].value = ws['D15'].value.replace('{supplier.documents.cnpj}', '{project.document}')
ws['D16'].value = ws['D16'].value.replace('{supplier.documents.cnpj}', '{project.ie}')
wb.save('.fractal/artifacts/docs/templates/PEDIDO-GERAL.xlsx')
print('ok')
